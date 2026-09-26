#!/usr/bin/env python3
"""
audit.py — универсальный аудит расходов из CSV / Excel файла.

Использование:
    python audit.py <путь_к_файлу>

Поддерживаются: .csv, .tsv, .txt (любой разделитель определяется автоматически)
                .xlsx, .xlsm (нужен пакет openpyxl: pip install openpyxl)

Столбцы определяются автоматически по названиям заголовков
(рус./англ. синонимы) — файл не обязан называться expenses.csv
и колонки могут идти в любом порядке.

Скрипт выводит отчёт в консоль и сохраняет его в report.txt:
  1. Общая сумма расходов (точная, без округлений)
  2. Сумма по каждой категории, по убыванию
  3. Топ-5 самых крупных трат
  4. Дубликаты (одинаковые дата + описание + сумма)
  5. Аномалии (отклонение от среднего по категории > 3 стандартных отклонений)
  6. Строки с отрицательными суммами
"""

import os
import sys
import csv
import statistics
from collections import defaultdict, Counter
from decimal import Decimal, InvalidOperation


# --- Автоопределение столбцов: роль -> возможные названия заголовков ---
COLUMN_SYNONYMS = {
    "date": [
        "date", "дата", "день", "число", "day", "period", "период", "when",
    ],
    "category": [
        "category", "категория", "cat", "тип", "вид", "группа", "group",
        "статья", "раздел", "type",
    ],
    "amount": [
        "amount", "сумма", "сум", "sum", "стоимость", "цена", "price",
        "total", "итого", "руб", "value", "расход", "затраты", "cost",
    ],
    "description": [
        "description", "описание", "наименование", "назначение", "comment",
        "комментарий", "примечание", "note", "details", "детали", "name",
        "контрагент", "получатель",
    ],
}


def normalize(text):
    return (text or "").strip().lower().replace("ё", "е")


def detect_columns(headers):
    """
    По списку заголовков подбираем, какой столбец за какую роль отвечает.
    Возвращает словарь role -> имя_заголовка (или None, если не найдено).
    """
    mapping = {"date": None, "category": None, "amount": None, "description": None}
    norm_headers = [(h, normalize(h)) for h in headers]

    for role, synonyms in COLUMN_SYNONYMS.items():
        # 1) точное совпадение заголовка с синонимом
        for original, norm in norm_headers:
            if norm in synonyms and original not in mapping.values():
                mapping[role] = original
                break
        if mapping[role]:
            continue
        # 2) заголовок содержит синоним как подстроку
        for original, norm in norm_headers:
            if original in mapping.values():
                continue
            if any(syn in norm for syn in synonyms):
                mapping[role] = original
                break
    return mapping


# --- Чтение таблицы из разных форматов ---
def read_table(path):
    """Возвращает (headers, rows) где rows — список dict {header: value}."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path)
    return _read_csv(path)


def _read_csv(path):
    with open(path, "r", newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel  # запасной вариант: обычная запятая
        reader = csv.DictReader(f, dialect=dialect)
        headers = reader.fieldnames or []
        rows = [dict(r) for r in reader]
    return headers, rows


def _read_excel(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit(
            "Для чтения Excel нужен пакет openpyxl.\n"
            "Установите его командой:  pip install openpyxl\n"
            "Либо сохраните файл как CSV и запустите снова."
        )
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [str(h) if h is not None else "" for h in header_row]
    rows = []
    for values in rows_iter:
        if values is None or all(v is None for v in values):
            continue
        row = {}
        for h, v in zip(headers, values):
            row[h] = "" if v is None else str(v)
        rows.append(row)
    return headers, rows


def parse_amount(raw):
    """Аккуратно разбираем сумму: убираем пробелы, валюту, приводим запятую к точке."""
    text = str(raw or "").strip()
    # оставляем только цифры, точку, запятую и минус
    cleaned = "".join(ch for ch in text if ch.isdigit() or ch in ".,-")
    cleaned = cleaned.replace(" ", "")
    # если есть и точка и запятая — считаем запятую разделителем тысяч
    if "," in cleaned and "." in cleaned:
        cleaned = cleaned.replace(",", "")
    else:
        cleaned = cleaned.replace(",", ".")
    return Decimal(cleaned)


def load_rows(path):
    """Читаем файл, определяем столбцы, возвращаем (rows, columns, skipped)."""
    headers, raw_rows = read_table(path)
    if not headers:
        sys.exit("Файл пустой или не удалось прочитать заголовки.")

    columns = detect_columns(headers)
    if not columns["amount"]:
        sys.exit(
            "Не удалось определить столбец с суммой.\n"
            f"Заголовки файла: {headers}\n"
            "Переименуйте столбец с суммой (например, в 'amount' или 'сумма')."
        )

    rows = []
    skipped = 0
    for i, raw in enumerate(raw_rows, start=2):  # строка 1 — заголовок
        try:
            amount = parse_amount(raw.get(columns["amount"]))
        except (InvalidOperation, AttributeError):
            skipped += 1
            continue
        rows.append(
            {
                "line": i,
                "date": (raw.get(columns["date"], "") or "").strip() if columns["date"] else "",
                "category": (raw.get(columns["category"], "") or "").strip() if columns["category"] else "",
                "description": (raw.get(columns["description"], "") or "").strip() if columns["description"] else "",
                "amount": amount,
            }
        )
    return rows, columns, skipped


def fmt(x):
    return f"{x:,.2f}".replace(",", " ")


def build_report(path, rows, columns, skipped):
    out = []
    w = out.append

    w("=" * 60)
    w("           ОТЧЁТ ПО АУДИТУ РАСХОДОВ")
    w("=" * 60)
    w(f"Файл:            {path}")
    w(f"Всего записей:   {len(rows)}")
    if skipped:
        w(f"Пропущено строк (битая сумма): {skipped}")
    detected = ", ".join(f"{role}={col}" for role, col in columns.items() if col)
    w(f"Определены столбцы: {detected}")
    w("")

    # 1. Общая сумма (точная, Decimal)
    total = sum((r["amount"] for r in rows), Decimal("0"))
    w("1) ОБЩАЯ СУММА РАСХОДОВ")
    w("-" * 60)
    w(f"   {fmt(total)}   (точно: {total})")
    w("")

    # 2. Сумма по категориям, по убыванию
    by_cat = defaultdict(lambda: Decimal("0"))
    cnt_cat = Counter()
    for r in rows:
        cat = r["category"] if r["category"] else "(без категории)"
        by_cat[cat] += r["amount"]
        cnt_cat[cat] += 1
    w("2) СУММА ПО КАТЕГОРИЯМ (по убыванию)")
    w("-" * 60)
    if not columns["category"]:
        w("   (столбец категории не найден — разбивка недоступна)")
    else:
        for cat, s in sorted(by_cat.items(), key=lambda kv: kv[1], reverse=True):
            share = (s / total * 100) if total else Decimal("0")
            w(f"   {fmt(s):>18}  {share:5.1f}%  ({cnt_cat[cat]:>4} шт)  {cat}")
    w("")

    # 3. Топ-5 самых крупных трат
    w("3) ТОП-5 САМЫХ КРУПНЫХ ТРАТ")
    w("-" * 60)
    for r in sorted(rows, key=lambda r: r["amount"], reverse=True)[:5]:
        parts = [p for p in (r["date"], r["category"], r["description"]) if p]
        w(f"   {fmt(r['amount']):>18}  {'  '.join(parts)}")
    w("")

    # 4. Дубликаты: одинаковые дата + описание + сумма
    seen = defaultdict(list)
    for r in rows:
        key = (r["date"], r["description"], r["amount"])
        seen[key].append(r)
    dups = {k: v for k, v in seen.items() if len(v) > 1}
    w("4) ДУБЛИКАТЫ (одинаковые дата + описание + сумма)")
    w("-" * 60)
    if not dups:
        w("   Дубликаты не найдены.")
    else:
        extra = sum(len(v) - 1 for v in dups.values())
        w(f"   Видов дубликатов: {len(dups)}, лишних копий: {extra}")
        for (date, desc, amount), group in sorted(dups.items(), key=lambda kv: kv[0][2], reverse=True):
            lines = ", ".join(str(r["line"]) for r in group)
            w(f"   x{len(group)}  {fmt(amount):>18}  {date}  {desc}   [строки: {lines}]")
    w("")

    # 5. Аномалии: |amount - среднее_по_категории| > 3 * std_по_категории
    cat_values = defaultdict(list)
    for r in rows:
        cat = r["category"] if r["category"] else "(без категории)"
        cat_values[cat].append(float(r["amount"]))
    thresholds = {}
    for cat, vals in cat_values.items():
        if len(vals) >= 2:
            mean = statistics.mean(vals)
            std = statistics.stdev(vals)
            thresholds[cat] = (mean, std)
    anomalies = []
    for r in rows:
        cat = r["category"] if r["category"] else "(без категории)"
        if cat in thresholds:
            mean, std = thresholds[cat]
            if std > 0 and abs(float(r["amount"]) - mean) > 3 * std:
                anomalies.append((r, mean, std))
    w("5) АНОМАЛИИ (отклонение от среднего по категории > 3 сигм)")
    w("-" * 60)
    if not anomalies:
        w("   Аномалии не найдены.")
    else:
        w(f"   Найдено: {len(anomalies)}")
        for r, mean, std in sorted(anomalies, key=lambda t: t[0]["amount"], reverse=True):
            dev = (float(r["amount"]) - mean) / std
            parts = [p for p in (r["date"], r["category"], r["description"]) if p]
            w(f"   {fmt(r['amount']):>18}  ({dev:+.1f} сигм)  {'  '.join(parts)}  [строка {r['line']}]")
    w("")

    # 6. Отрицательные суммы
    negatives = [r for r in rows if r["amount"] < 0]
    w("6) СТРОКИ С ОТРИЦАТЕЛЬНЫМИ СУММАМИ")
    w("-" * 60)
    if not negatives:
        w("   Отрицательных сумм не найдено.")
    else:
        w(f"   Найдено: {len(negatives)}")
        for r in sorted(negatives, key=lambda r: r["amount"]):
            parts = [p for p in (r["date"], r["category"], r["description"]) if p]
            w(f"   {fmt(r['amount']):>18}  {'  '.join(parts)}  [строка {r['line']}]")
    w("")

    w("=" * 60)
    w("Конец отчёта.")
    return "\n".join(out)


def main():
    # Windows-консоль по умолчанию cp1251 и не умеет печатать некоторые символы —
    # переключаем стандартный вывод на UTF-8, чтобы отчёт печатался корректно.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass

    if len(sys.argv) != 2:
        sys.exit("Использование: python audit.py <путь_к_файлу>")
    path = sys.argv[1]
    if not os.path.exists(path):
        sys.exit(f"Ошибка: файл не найден — {path}")

    rows, columns, skipped = load_rows(path)
    if not rows:
        sys.exit("В файле нет данных для анализа.")

    report = build_report(path, rows, columns, skipped)
    print(report)

    with open("report.txt", "w", encoding="utf-8") as f:
        f.write(report + "\n")
    print("\n[OK] Отчёт сохранён в report.txt")


if __name__ == "__main__":
    main()
